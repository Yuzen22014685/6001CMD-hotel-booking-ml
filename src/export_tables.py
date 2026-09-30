"""
Export every CSV in /tables into a single formatted Word document
(tables/report_tables.docx) for copying into the report.
Run from the project root:  python src/export_tables.py
"""
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt, Cm

ROOT = Path(__file__).resolve().parent.parent
TAB_DIR = ROOT / "tables"
OUT_FILE = TAB_DIR / "report_tables.docx"

# Report titles for each table (add new entries as tables are created)
TITLES = {
    "table1.1_feature_overview": "Overview of Dataset Features, Types, Missing Values and Descriptions",
    "table1.2_cancellation_by_hotel": "Cancellation Rate by Hotel Type",
    "table1.3_challenge_scan": "Dataset Challenge Scan",
}

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(9)


def table_sort_key(path):
    """Sort files by table number, e.g. table1.2 before table1.10."""
    match = re.match(r"table(\d+)\.(\d+)", path.stem)
    return (int(match.group(1)), int(match.group(2))) if match else (999, 999)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text))
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    run.bold = bold


def main():
    csv_files = sorted(TAB_DIR.glob("table*.csv"), key=table_sort_key)
    if not csv_files:
        print("No table CSV files found in /tables.")
        return

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)

    for path in csv_files:
        # Read everything as text so values appear exactly as saved
        data = pd.read_csv(path, dtype=str, keep_default_na=False)
        match = re.match(r"table(\d+\.\d+)", path.stem)
        number = match.group(1) if match else path.stem
        title = TITLES.get(path.stem, path.stem.replace("_", " "))

        # APA 7 style: bold table number, italic title on the next line
        p_num = doc.add_paragraph()
        r_num = p_num.add_run(f"Table {number}")
        r_num.bold = True
        r_num.font.name = FONT_NAME
        p_title = doc.add_paragraph()
        r_title = p_title.add_run(title)
        r_title.italic = True
        r_title.font.name = FONT_NAME

        table = doc.add_table(rows=1, cols=len(data.columns))
        table.style = "Table Grid"
        for i, col in enumerate(data.columns):
            set_cell_text(table.rows[0].cells[i], col, bold=True)
        for _, row in data.iterrows():
            cells = table.add_row().cells
            for i, value in enumerate(row):
                set_cell_text(cells[i], value)

        doc.add_paragraph()
        print(f"Added Table {number}: {title} ({len(data)} rows)")

    try:
        doc.save(OUT_FILE)
        print(f"\nSaved: {OUT_FILE}")
    except PermissionError:
        print("\nCould not save. Close report_tables.docx in Word and run again.")


if __name__ == "__main__":
    main()